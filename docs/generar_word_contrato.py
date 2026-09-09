"""Genera el Word editable del contrato ModoOps v2 desde su fuente SSOT.

SSOT: docs/contrato-modoops-completo.html -> docs/contrato-modoops.docx
(editable en Word/LibreOffice; el PDF se regenera con docs/generar_pdf_contrato.py)

Uso:  python docs/generar_word_contrato.py [--check]
--check: solo verifica que la fuente existe (para CI, fail-closed).
"""
from __future__ import annotations

import sys
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FORMAL_HTML = ROOT / "docs" / "contrato-modoops-completo.html"
WORD_OUT = ROOT / "docs" / "contrato-modoops.docx"


class ContractDocxParser(HTMLParser):
    """Convierte el HTML del contrato formal a un documento python-docx."""

    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self._buf = ""
        self._bold = False
        self._runs = []  # (texto, bold)
        self._mode = "body"  # h1|h2|h3|body|bullet|sub|small
        self._skip = 0
        self._table = None
        self._row = None
        self._cells = None

    # -- helpers ---------------------------------------------------------
    def _push(self, text: str):
        if text:
            self._runs.append((text, self._bold))

    def _flush(self):
        runs = [(t, b) for t, b in self._runs if t.strip()]
        self._runs = []
        self._buf = ""
        if not runs:
            return
        if self._cells is not None:
            self._cells.append(runs)
            return
        if self._mode == "bullet":
            para = self.doc.add_paragraph(style="List Bullet")
        elif self._mode == "h1":
            para = self.doc.add_heading(level=1)
            para.alignment = 1
        elif self._mode == "h2":
            para = self.doc.add_heading(level=2)
        elif self._mode == "h3":
            para = self.doc.add_heading(level=3)
        else:
            para = self.doc.add_paragraph()
            if self._mode == "sub":
                para.alignment = 1
        for text, bold in runs:
            run = para.add_run(text)
            run.bold = bold

    # -- parser ----------------------------------------------------------
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        cls = attrs.get("class", "")
        if tag in ("head", "style", "script"):
            self._skip += 1
            return
        if self._skip:
            return
        if tag == "div" and cls == "firma":
            self._flush()
            self._cells = []
            return
        if tag == "div" and self._cells is not None:
            self._flush()
            return
        if tag == "div" and cls == "pagebreak":
            self._flush()
            self.doc.add_page_break()
            return
        if tag in ("h1", "h2", "h3", "p"):
            self._flush()
            if tag == "p" and cls in ("sub", "small"):
                self._mode = "sub"
            else:
                self._mode = tag if tag in ("h1", "h2", "h3") else "body"
        elif tag == "li":
            self._flush()
            self._mode = "bullet"
        elif tag in ("ul", "ol"):
            self._flush()
        elif tag == "br":
            self._push("\n")
        elif tag in ("b", "strong"):
            if self._buf:
                self._push(self._buf)
                self._buf = ""
            self._bold = True
        elif tag == "table":
            self._flush()
            self._table = []
        elif tag == "tr":
            self._row = []
        elif tag == "td":
            self._buf = ""

    def handle_endtag(self, tag):
        if tag in ("head", "style", "script"):
            self._skip = max(0, self._skip - 1)
            return
        if self._skip:
            return
        if tag in ("b", "strong"):
            if self._buf:
                self._push(self._buf)
                self._buf = ""
            self._bold = False
            return
        if tag == "div" and self._cells is not None:
            self._flush()
            cells = self._cells
            self._cells = None
            while len(cells) < 2:
                cells.append([])
            table = self.doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            for i, runs in enumerate(cells[:2]):
                cell = table.cell(0, i)
                for j, (text, bold) in enumerate(runs):
                    para = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                    run = para.add_run(text)
                    run.bold = bold
            self.doc.add_paragraph()
            return
        if tag in ("h1", "h2", "h3", "p", "li"):
            if self._buf:
                self._push(self._buf)
                self._buf = ""
            self._flush()
            self._mode = "body"
        elif tag == "td":
            self._row.append(self._buf.strip())
            self._buf = ""
        elif tag == "tr" and self._row is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            table = self.doc.add_table(rows=len(self._table), cols=len(self._table[0]))
            table.style = "Table Grid"
            for r, row in enumerate(self._table):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
            self.doc.add_paragraph()
            self._table = None
        elif tag == "div":
            self._flush()

    def handle_data(self, data):
        if self._skip:
            return
        if self._cells is not None and not self._runs and not self._buf:
            # espacios entre celdas de firma: se ignoran al hacer flush
            if not data.strip():
                return
        self._buf += data


def main() -> int:
    if not FORMAL_HTML.exists():
        print(f"Falta fuente SSOT: {FORMAL_HTML}")
        return 1
    if "--check" in sys.argv:
        print("Fuente Word OK.")
        return 0
    from docx import Document

    doc = Document()
    parser = ContractDocxParser(doc)
    parser.feed(FORMAL_HTML.read_text(encoding="utf-8"))
    doc.save(str(WORD_OUT))
    print(f"Word guardado en: {WORD_OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
